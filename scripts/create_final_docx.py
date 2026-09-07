from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# 设置默认字体
style = doc.styles['Normal']
font = style.font
font.name = 'Microsoft YaHei'
font.size = Pt(11)

# 标题
title = doc.add_heading('个性化智能营养助手', 0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle = doc.add_paragraph('项目需求说明与知识库建设（重构版）')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph('')

# 一、项目概述
doc.add_heading('一、项目概述', level=1)

doc.add_heading('1.1 核心理念', level=2)
doc.add_paragraph('"越用越懂你" - 通过每次对话不断充实用户画像，形成专属知识库')

doc.add_heading('1.2 系统定位', level=2)
doc.add_paragraph('本系统是智能饮食推荐助手，不是医疗诊断系统。')
doc.add_paragraph('• 提供饮食建议和营养参考')
doc.add_paragraph('• 不能替代医生、营养师的专业诊疗')
doc.add_paragraph('• 不能用于疾病诊断、治疗、预防')

doc.add_heading('1.3 目标用户', level=2)
users = [
    '减脂人群：想要减重、降低体脂率',
    '增肌人群：健身增肌、提升肌肉量',
    '健康管理人群：慢病管理（糖尿病、高血压等）',
    '亚健康人群：睡眠差、压力大、易疲劳',
    '健身爱好者：追求健康生活方式',
]
for user in users:
    doc.add_paragraph(user, style='List Bullet')

doc.add_heading('1.4 核心功能', level=2)
features = [
    '动态用户画像：通过对话自动采集和更新用户信息',
    '智能营养计算：基于BMR/TDEE计算每日营养需求',
    '个性化推荐：根据用户画像和偏好推荐菜谱',
    '热量追踪：记录每餐摄入，实时监控热量',
    '偏好学习：AI自动学习用户口味和饮食习惯',
    '周计划生成：自动生成一周饮食计划',
    '安全风控：危险饮食识别、过敏校验、医疗边界',
]
for feature in features:
    doc.add_paragraph(feature, style='List Bullet')

# 二、安全与健康风险控制
doc.add_heading('二、安全与健康风险控制', level=1)

doc.add_heading('2.1 医疗边界声明', level=2)
doc.add_paragraph('⚠️ 健康提示')
doc.add_paragraph('本系统提供的饮食建议仅供参考，不构成医疗建议。')
doc.add_paragraph('如有健康问题，请咨询专业医生或注册营养师。')

doc.add_paragraph('')
doc.add_paragraph('特殊情况处理：')
table = doc.add_table(rows=6, cols=2)
table.style = 'Table Grid'
headers = ['人群', '处理方式']
for i, header in enumerate(headers):
    table.rows[0].cells[i].text = header
data = [
    ['糖尿病患者', '提示"请遵医嘱，本系统仅供参考"'],
    ['高血压患者', '提示"请遵医嘱，注意钠摄入"'],
    ['肾病患者', '提示"请遵医嘱，注意蛋白质摄入"'],
    ['孕妇/哺乳期', '提示"请遵医嘱，注意额外营养需求"'],
    ['严重过敏者', '提示"请仔细核实食材，如有疑问请咨询医生"'],
]
for i, row_data in enumerate(data):
    for j, cell_data in enumerate(row_data):
        table.rows[i+1].cells[j].text = cell_data

doc.add_heading('2.2 过敏风险校验', level=2)
doc.add_paragraph('过敏校验流程：')
doc.add_paragraph('1. 菜谱过敏原标签完整 → 检查与用户过敏原是否冲突')
doc.add_paragraph('2. 菜谱过敏原标签缺失 → 标记为"过敏原未知"，不推荐给高敏用户')

doc.add_paragraph('')
doc.add_paragraph('过敏风险等级：')
table = doc.add_table(rows=4, cols=3)
table.style = 'Table Grid'
headers = ['等级', '说明', '处理方式']
for i, header in enumerate(headers):
    table.rows[0].cells[i].text = header
data = [
    ['高风险', '菜谱明确含有用户过敏原', '强制过滤，不推荐'],
    ['中风险', '菜谱过敏原信息缺失', '标记警告，不推荐给高敏用户'],
    ['低风险', '菜谱不含用户过敏原', '正常推荐'],
]
for i, row_data in enumerate(data):
    for j, cell_data in enumerate(row_data):
        table.rows[i+1].cells[j].text = cell_data

doc.add_heading('2.3 异常输入处理', level=2)
doc.add_paragraph('危险饮食识别：')
doc.add_paragraph('• 极端低热量：只吃500大卡、一天只吃xxx')
doc.add_paragraph('• 极端断食：节食3天、断食5天')
doc.add_paragraph('• 饮食障碍：催吐、暴食')
doc.add_paragraph('• 危险饮食：只吃水果、零热量')

doc.add_paragraph('')
doc.add_paragraph('安全阈值：')
doc.add_paragraph('• 女性：不低于1200kcal/天')
doc.add_paragraph('• 男性：不低于1500kcal/天')
doc.add_paragraph('• 任何情况下：不低于基础代谢的80%')

doc.add_heading('2.4 用户隐私保护', level=2)
doc.add_paragraph('隐私保护措施：')
doc.add_paragraph('• 敏感数据AES-256加密存储')
doc.add_paragraph('• 严格的访问控制')
doc.add_paragraph('• 聊天记录保留90天，饮食记录保留1年')
doc.add_paragraph('• 用户可随时删除所有数据')

# 三、用户画像与动态学习机制
doc.add_heading('三、用户画像与动态学习机制', level=1)

doc.add_heading('3.1 置信度机制', level=2)
doc.add_paragraph('为每条提取信息增加置信度分数（0-1）：')

table = doc.add_table(rows=5, cols=3)
table.style = 'Table Grid'
headers = ['置信度范围', '说明', '来源示例']
for i, header in enumerate(headers):
    table.rows[0].cells[i].text = header
data = [
    ['0.9-1.0', '高置信度', '用户明确陈述'],
    ['0.7-0.9', '中置信度', '用户间接表达'],
    ['0.5-0.7', '低置信度', '模糊表述、推断'],
    ['0.3-0.5', '极低置信度', '猜测、不确定'],
]
for i, row_data in enumerate(data):
    for j, cell_data in enumerate(row_data):
        table.rows[i+1].cells[j].text = cell_data

doc.add_heading('3.2 矛盾信息处理', level=2)
doc.add_paragraph('当新旧信息冲突时：')
doc.add_paragraph('1. 置信度相近 → AI主动确认："你之前说A，现在说B，哪个正确？"')
doc.add_paragraph('2. 新信息置信度更高 → 覆盖旧信息')
doc.add_paragraph('3. 旧信息置信度更高 → 保留旧信息')

doc.add_heading('3.3 信息缺失检测', level=2)
doc.add_paragraph('字段分类：')
table = doc.add_table(rows=8, cols=3)
table.style = 'Table Grid'
headers = ['字段', '类型', '说明']
for i, header in enumerate(headers):
    table.rows[0].cells[i].text = header
data = [
    ['性别、年龄、身高、体重', '必填', '影响BMR计算'],
    ['慢性疾病、食物过敏', '必填', '影响饮食安全'],
    ['活动水平', '建议填写', '影响TDEE计算'],
    ['口味偏好', '可选', '提升推荐满意度'],
    ['烹饪技能', '可选', '推荐难度匹配'],
    ['预算', '可选', '成本控制'],
    ['运动类型', '可选', '个性化推荐'],
]
for i, row_data in enumerate(data):
    for j, cell_data in enumerate(row_data):
        table.rows[i+1].cells[j].text = cell_data

doc.add_heading('3.4 行为反馈细粒度设计', level=2)
doc.add_paragraph('拒绝菜品时提供反馈标签：')
doc.add_paragraph('• 口味相关：太辣、太甜、太咸、不好吃')
doc.add_paragraph('• 健康相关：热量太高、不够饱、太腻')
doc.add_paragraph('• 实用性相关：太难做、食材买不到、太贵、太耗时')
doc.add_paragraph('• 临时原因：今天不想吃、已经吃过了、不饿')

doc.add_paragraph('')
doc.add_paragraph('黑名单管理：')
doc.add_paragraph('• 单次拒绝不进入黑名单')
doc.add_paragraph('• 累计3次相同拒绝才标记为"不喜欢"')
doc.add_paragraph('• 区分临时拒绝和长期忌口')

# 四、知识库与RAG设计
doc.add_heading('四、知识库与RAG设计', level=1)

doc.add_heading('4.1 动态分块策略', level=2)
doc.add_paragraph('分块原则：')
doc.add_paragraph('• 语义完整性：每个块包含完整语义单元')
doc.add_paragraph('• 长度灵活性：不再硬性限制500-800字符')
doc.add_paragraph('• 重叠合理性：按段落/语义重叠，不按字符')

doc.add_paragraph('')
doc.add_paragraph('菜谱分块类型：')
doc.add_paragraph('1. 基础信息块（菜名、分类、热量）')
doc.add_paragraph('2. 食材块（食材列表、过敏原）')
doc.add_paragraph('3. 步骤块（烹饪步骤，长菜谱拆分）')
doc.add_paragraph('4. 营养块（营养成分）')

doc.add_heading('4.2 公私库分离', level=2)
doc.add_paragraph('知识库架构：')
doc.add_paragraph('• 公共向量库：菜谱、营养知识（全局共享）')
doc.add_paragraph('• 用户私有库：用户画像、饮食记录、聊天洞察')

doc.add_paragraph('')
doc.add_paragraph('检索流程：')
doc.add_paragraph('1. 公共菜谱库向量检索（Top-20）')
doc.add_paragraph('2. 用户私有库过滤（排除过敏原、忌口）')
doc.add_paragraph('3. 用户偏好重排（口味、热量、难度）')
doc.add_paragraph('4. 最终推荐（Top-5）')

doc.add_heading('4.3 质量过滤机制', level=2)
doc.add_paragraph('数据可信度分级：')
table = doc.add_table(rows=5, cols=3)
table.style = 'Table Grid'
headers = ['级别', '来源', '处理方式']
for i, header in enumerate(headers):
    table.rows[0].cells[i].text = header
data = [
    ['官方数据', '中国食物成分表、USDA', '直接使用'],
    ['用户验证', '点赞>100、差评<5%', '正常使用'],
    ['用户生成', '普通UGC内容', '标记后使用'],
    ['可疑数据', '营养值异常', '过滤不使用'],
]
for i, row_data in enumerate(data):
    for j, cell_data in enumerate(row_data):
        table.rows[i+1].cells[j].text = cell_data

doc.add_heading('4.4 元数据补全', level=2)
doc.add_paragraph('新增菜谱元数据：')
doc.add_paragraph('• cooking_time: 烹饪时间')
doc.add_paragraph('• difficulty: 难度等级（1-5）')
doc.add_paragraph('• ingredient_availability: 食材易得性')
doc.add_paragraph('• is_takeout_friendly: 是否适合外卖替代')
doc.add_paragraph('• is_vegetarian: 是否素食')
doc.add_paragraph('• cost_level: 成本等级（1-5）')

doc.add_paragraph('')
doc.add_paragraph('新增用户画像字段：')
doc.add_paragraph('• exercise_types: 运动类型')
doc.add_paragraph('• dietary_restrictions: 饮食禁忌')
doc.add_paragraph('• current_medications: 正在服用的药物')
doc.add_paragraph('• pregnancy_status: 孕期状态')

# 五、推荐算法设计
doc.add_heading('五、推荐算法设计', level=1)

doc.add_heading('5.1 特殊人群营养计算', level=2)
doc.add_paragraph('特殊人群修正：')
table = doc.add_table(rows=6, cols=3)
table.style = 'Table Grid'
headers = ['人群', '热量调整', '特殊限制']
for i, header in enumerate(headers):
    table.rows[0].cells[i].text = header
data = [
    ['孕妇', '+300-500kcal', '注意叶酸、铁'],
    ['哺乳期', '+500kcal', '注意钙、DHA'],
    ['老年人（>65岁）', 'BMR×0.9', '蛋白质适量'],
    ['糖尿病患者', '正常', '碳水≤45%'],
    ['肾病患者', '正常', '蛋白质0.6-0.8g/kg'],
]
for i, row_data in enumerate(data):
    for j, cell_data in enumerate(row_data):
        table.rows[i+1].cells[j].text = cell_data

doc.add_heading('5.2 约束冲突降级', level=2)
doc.add_paragraph('约束优先级：')
doc.add_paragraph('1. 关键约束（过敏、慢病）- 必须满足')
doc.add_paragraph('2. 高优先级（目标、热量）- 必须满足')
doc.add_paragraph('3. 中优先级（口味、偏好）- 可放宽')
doc.add_paragraph('4. 低优先级（其他）- 可放宽')

doc.add_paragraph('')
doc.add_paragraph('降级逻辑：')
doc.add_paragraph('• 无完全匹配时，按优先级逐步放宽约束')
doc.add_paragraph('• 向用户说明："没有完全匹配，为你适度放宽XX条件"')

doc.add_heading('5.3 周饮食计划生成', level=2)
doc.add_paragraph('生成原则：')
doc.add_paragraph('• 每周营养均衡（总热量、蛋白质、碳水、脂肪）')
doc.add_paragraph('• 食材重复度控制（同类食材每周≤3次）')
doc.add_paragraph('• 三餐轮换（早餐/午餐/晚餐类型多样化）')

# 六、工程落地
doc.add_heading('六、工程落地', level=1)

doc.add_heading('6.1 评估指标体系', level=2)
doc.add_paragraph('核心指标：')
table = doc.add_table(rows=7, cols=3)
table.style = 'Table Grid'
headers = ['指标', '说明', '目标']
for i, header in enumerate(headers):
    table.rows[0].cells[i].text = header
data = [
    ['推荐接受率', '用户接受推荐菜品的比例', '>70%'],
    ['推荐拒绝率', '用户拒绝推荐菜品的比例', '<15%'],
    ['热量达标率', '每日热量摄入在目标±10%范围内', '>85%'],
    ['蛋白质达标率', '每日蛋白质摄入达标', '>80%'],
    ['满意度评分', '用户对推荐的平均评分（1-5分）', '>4.0'],
    ['留存率', '用户继续使用的比例', '>60%'],
]
for i, row_data in enumerate(data):
    for j, cell_data in enumerate(row_data):
        table.rows[i+1].cells[j].text = cell_data

doc.add_heading('6.2 知识库更新机制', level=2)
doc.add_paragraph('更新策略：')
doc.add_paragraph('• 增量更新：新菜谱直接追加')
doc.add_paragraph('• 全量重建：营养数据变更时重建索引')
doc.add_paragraph('• 版本控制：保留历史版本，支持回滚')
doc.add_paragraph('• 更新日志：记录每次更新')

doc.add_heading('6.3 部署架构', level=2)
doc.add_paragraph('系统架构：')
doc.add_paragraph('• 负载均衡器（Nginx）')
doc.add_paragraph('• 应用服务器集群（FastAPI）')
doc.add_paragraph('• 数据层：Redis（缓存）+ PostgreSQL（用户数据）+ FAISS（向量索引）')

# 七、实施计划
doc.add_heading('七、实施计划', level=1)

table = doc.add_table(rows=5, cols=4)
table.style = 'Table Grid'
headers = ['阶段', '任务', '时间', '交付物']
for i, header in enumerate(headers):
    table.rows[0].cells[i].text = header
data = [
    ['第一阶段', '安全与核心功能', '2-3天', '医疗边界声明、过敏校验、异常处理、信息缺失检测'],
    ['第二阶段', '知识库重构', '2-3天', '动态分块、公私库分离、质量过滤、元数据补全'],
    ['第三阶段', '推荐算法优化', '1-2天', '特殊人群计算、约束降级、周计划生成'],
    ['第四阶段', '工程与评估', '1-2天', '评估指标、更新机制、隐私保护'],
]
for i, row_data in enumerate(data):
    for j, cell_data in enumerate(row_data):
        table.rows[i+1].cells[j].text = cell_data

doc.add_paragraph('')
doc.add_paragraph('总计：6-10天')

# 保存文档
output_path = r'C:\Users\Sean.bx\Desktop\project_based_training\nutri_assistant\docs\reports\项目需求说明与知识库建设_重构版.docx'
doc.save(output_path)
print(f"文档已保存到: {output_path}")
